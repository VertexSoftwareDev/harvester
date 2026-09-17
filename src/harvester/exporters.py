"""Exporters write scraped items to files.

Every exporter is a class with an ``extension`` and an ``export(items, path)``
method. Register new ones with ``@exporter_registry.register("name")`` or through
the ``harvester.exporters`` entry-point group.
"""

import csv
import json
from collections.abc import Sequence
from pathlib import Path
from typing import Any, ClassVar, Protocol

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from harvester.item import Item
from harvester.plugins import Registry


class Exporter(Protocol):
    extension: ClassVar[str]

    def export(self, items: Sequence[Item], path: Path) -> None: ...


exporter_registry: Registry[type[Exporter]] = Registry(
    "exporter", entry_point_group="harvester.exporters"
)


def columns_of(items: Sequence[Item]) -> list[str]:
    """Union of all keys, in first-seen order (items may have different shapes)."""
    return list(dict.fromkeys(key for item in items for key in item))


def _flatten(value: Any) -> Any:
    """Make nested values fit into a single spreadsheet cell."""
    if isinstance(value, list | tuple | set):
        return ", ".join(str(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return value


@exporter_registry.register("csv")
class CsvExporter:
    """CSV with a UTF-8 BOM, so Excel opens non-ASCII text correctly."""

    extension: ClassVar[str] = "csv"

    def export(self, items: Sequence[Item], path: Path) -> None:
        columns = columns_of(items)
        with path.open("w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.DictWriter(fh, fieldnames=columns, restval="")
            writer.writeheader()
            for item in items:
                writer.writerow({k: "" if v is None else _flatten(v) for k, v in item.items()})


@exporter_registry.register("json")
class JsonExporter:
    """A single pretty-printed JSON array."""

    extension: ClassVar[str] = "json"

    def export(self, items: Sequence[Item], path: Path) -> None:
        with path.open("w", encoding="utf-8") as fh:
            json.dump(list(items), fh, ensure_ascii=False, indent=2, default=str)
            fh.write("\n")


@exporter_registry.register("jsonl")
class JsonLinesExporter:
    """One JSON object per line; easy to stream and append."""

    extension: ClassVar[str] = "jsonl"

    def export(self, items: Sequence[Item], path: Path) -> None:
        with path.open("w", encoding="utf-8") as fh:
            for item in items:
                fh.write(json.dumps(item, ensure_ascii=False, default=str) + "\n")


@exporter_registry.register("excel")
class ExcelExporter:
    """Excel workbook with a styled, filterable header, sized columns and clickable links."""

    extension: ClassVar[str] = "xlsx"
    max_column_width: ClassVar[int] = 60

    def export(self, items: Sequence[Item], path: Path) -> None:
        columns = columns_of(items)
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Items"

        sheet.append(columns)
        for item in items:
            sheet.append([self._cell_value(item.get(column)) for column in columns])

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="2F5597")
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        sheet.freeze_panes = "A2"
        if columns:
            sheet.auto_filter.ref = sheet.dimensions

        for index, column_cells in enumerate(sheet.iter_cols(), start=1):
            width = max((len(str(c.value)) for c in column_cells if c.value is not None), default=8)
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = min(width + 2, self.max_column_width)
            for cell in column_cells[1:]:
                if isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"

        workbook.save(path)

    @staticmethod
    def _cell_value(value: Any) -> Any:
        value = _flatten(value)
        if isinstance(value, str):
            # Control characters in scraped text make openpyxl refuse to save.
            return ILLEGAL_CHARACTERS_RE.sub("", value)
        return value
