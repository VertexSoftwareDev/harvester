import csv
import json
from datetime import datetime
from pathlib import Path

import httpx
import respx
from openpyxl import load_workbook

from harvester.config import ExportSpec
from harvester.pipeline import Crawler, run_job
from tests.conftest import BASE, PAGE_1, PAGE_2, JobFactory


def mock_catalogue(respx_mock: respx.MockRouter) -> tuple[respx.Route, respx.Route]:
    return (
        respx_mock.get(BASE + "page-1.html").respond(200, html=PAGE_1),
        respx_mock.get(BASE + "page-2.html").respond(200, html=PAGE_2),
    )


async def test_crawl_follows_pagination_without_revisiting(
    respx_mock: respx.MockRouter, make_job: JobFactory
) -> None:
    page_1, page_2 = mock_catalogue(respx_mock)
    crawler = Crawler(make_job())

    titles = [item["title"] async for item in crawler.crawl()]

    assert titles == ["A Light in the Attic", "Tipping the Velvet", "Soumission"]
    assert page_1.call_count == 1  # page 2 links back to page 1
    assert page_2.call_count == 1
    stats = crawler.stats
    assert (stats.pages_ok, stats.pages_failed, stats.items, stats.items_skipped) == (2, 0, 3, 1)
    assert stats.finished_at is not None


async def test_max_pages_limits_the_crawl(
    respx_mock: respx.MockRouter, make_job: JobFactory
) -> None:
    _, page_2 = mock_catalogue(respx_mock)
    crawler = Crawler(make_job(pagination={"next": "li.next a", "max_pages": 1}))

    items = [item async for item in crawler.crawl()]

    assert len(items) == 2
    assert not page_2.called


async def test_retries_transient_failures(
    respx_mock: respx.MockRouter, make_job: JobFactory
) -> None:
    route = respx_mock.get(BASE + "page-1.html").mock(
        side_effect=[
            httpx.ConnectError("connection reset"),
            httpx.Response(503),
            httpx.Response(200, html=PAGE_2),
        ]
    )
    job = make_job(pagination=None, http={"rate_limit": None, "backoff": 0, "retries": 3})
    crawler = Crawler(job)

    items = [item async for item in crawler.crawl()]

    assert [item["title"] for item in items] == ["Soumission"]
    assert route.call_count == 3


async def test_failed_pages_do_not_stop_the_crawl(
    respx_mock: respx.MockRouter, make_job: JobFactory
) -> None:
    respx_mock.get(BASE + "missing.html").respond(404)
    respx_mock.get(BASE + "page-2.html").respond(200, html=PAGE_2)
    job = make_job(start_urls=[BASE + "missing.html", BASE + "page-2.html"], pagination=None)
    crawler = Crawler(job)

    items = [item async for item in crawler.crawl()]

    assert len(items) == 1
    assert (crawler.stats.pages_ok, crawler.stats.pages_failed) == (1, 1)


async def test_client_errors_are_not_retried(
    respx_mock: respx.MockRouter, make_job: JobFactory
) -> None:
    route = respx_mock.get(BASE + "page-1.html").respond(403)
    crawler = Crawler(make_job())

    assert [item async for item in crawler.crawl()] == []
    assert route.call_count == 1


async def test_sends_configured_headers(respx_mock: respx.MockRouter, make_job: JobFactory) -> None:
    route = respx_mock.get(BASE + "page-2.html").respond(200, html=PAGE_2)
    job = make_job(
        start_urls=[BASE + "page-2.html"],
        pagination=None,
        http={"rate_limit": None, "headers": {"User-Agent": "test-agent/1.0"}},
    )

    [item async for item in Crawler(job).crawl()]

    assert route.calls.last.request.headers["User-Agent"] == "test-agent/1.0"


async def test_run_job_writes_every_export(
    respx_mock: respx.MockRouter, make_job: JobFactory, tmp_path: Path
) -> None:
    mock_catalogue(respx_mock)
    exports = [
        ExportSpec(format="csv", path="{name}-{date}.csv"),
        ExportSpec(format="json"),
        ExportSpec(format="jsonl", path="nested/dir/{name}.jsonl"),
        ExportSpec(format="excel"),
    ]

    report = await run_job(
        make_job(), exports=exports, output_dir=tmp_path, now=datetime(2026, 9, 17)
    )

    assert report.outputs == (
        tmp_path / "books-2026-09-17.csv",
        tmp_path / "books.json",
        tmp_path / "nested/dir/books.jsonl",
        tmp_path / "books.xlsx",
    )
    with report.outputs[0].open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert [row["price"] for row in rows] == ["51.77", "53.74", "50.1"]

    data = json.loads(report.outputs[1].read_text(encoding="utf-8"))
    assert data[2]["title"] == "Soumission"

    lines = report.outputs[2].read_text(encoding="utf-8").splitlines()
    assert len(lines) == 3

    sheet = load_workbook(report.outputs[3]).active
    assert sheet is not None
    assert sheet.max_row == 4


async def test_run_job_defaults_to_json(
    respx_mock: respx.MockRouter, make_job: JobFactory, tmp_path: Path
) -> None:
    mock_catalogue(respx_mock)

    report = await run_job(make_job(), output_dir=tmp_path)

    assert report.outputs == (tmp_path / "books.json",)
